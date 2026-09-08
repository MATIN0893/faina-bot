import asyncio
import io
import json
import os
import re
import time
from typing import Any
from urllib.parse import urlparse

from aiohttp import web
from aiogram import Bot, Dispatcher, F, types
from aiogram.filters import Command, CommandStart
from aiogram.types import BufferedInputFile
from aiogram.webhook.aiohttp_server import SimpleRequestHandler, setup_application
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from google import genai
from google.genai import types as genai_types
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from price_catalog import GoogleSheetsPriceCatalog

BOT_TOKEN = os.getenv("BOT_TOKEN")
GEMINI_KEY = os.getenv("GEMINI_KEY")
BASE_URL = (os.getenv("RENDER_EXTERNAL_URL") or os.getenv("BASE_URL") or "").rstrip("/")
WEBHOOK_PATH = "/webhook"
WEBHOOK_SECRET = os.getenv("WEBHOOK_SECRET")
MAX_TEXT = 12000

configured_models = os.getenv("GEMINI_MODELS", "gemini-2.5-flash,gemini-2.0-flash")
MODELS = tuple(model.strip() for model in configured_models.split(",") if model.strip())
if not MODELS:
    MODELS = ("gemini-2.5-flash",)

if not BOT_TOKEN:
    raise RuntimeError("BOT_TOKEN is not set")
if not GEMINI_KEY:
    raise RuntimeError("GEMINI_KEY is not set")

client = genai.Client(
    api_key=GEMINI_KEY,
    http_options=genai_types.HttpOptions(
        timeout=30000,
        retry_options=genai_types.HttpRetryOptions(attempts=1),
    ),
)

bot = Bot(token=BOT_TOKEN)
dp = Dispatcher()
chat_locks: dict[int, asyncio.Lock] = {}
runtime_task = None


def lock_for(chat_id: int) -> asyncio.Lock:
    return chat_locks.setdefault(chat_id, asyncio.Lock())


def clean_json(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.I)
    text = re.sub(r"\s*```$", "", text)
    return text.strip()


def normalize_data(data: dict[str, Any]) -> dict[str, Any]:
    title = str(data.get("title") or "Документ").strip()[:120] or "Документ"
    text = str(data.get("document_text") or "").strip()
    headers = data.get("headers") or []
    rows = data.get("rows") or []
    if not isinstance(headers, list):
        headers = []
    headers = [str(x) for x in headers]
    clean_rows = []
    if headers and isinstance(rows, list):
        for row in rows:
            if isinstance(row, list):
                clean_rows.append([str(row[i]) if i < len(row) and row[i] is not None else "" for i in range(len(headers))])
    return {"title": title, "language": data.get("language", "mixed"), "document_text": text, "headers": headers, "rows": clean_rows}


async def ask_gemini(contents: Any, system_prompt: str = "") -> str | None:
    payload = contents
    if system_prompt:
        payload = [system_prompt, contents] if not isinstance(contents, list) else [system_prompt, *contents]
    for model in MODELS:
        started = time.monotonic()
        try:
            response = await client.aio.models.generate_content(model=model, contents=payload)
            text = getattr(response, "text", None)
            if text:
                print(f"Gemini OK model={model} seconds={time.monotonic()-started:.2f}", flush=True)
                return text.strip()
            print(f"Gemini empty model={model}", flush=True)
        except Exception as e:
            print(f"Gemini failed model={model} seconds={time.monotonic()-started:.2f} error={type(e).__name__}: {e}", flush=True)
    return None


SYSTEM_PROMPT = (
    "Ты Фаина — полезный AI-помощник. Отвечай на языке пользователя. "
    "Поддерживай русский и таджикский, включая таджикские буквы Ғ Ӣ Қ Ӯ Ҳ Ҷ. "
    "Не выдумывай данные. Не называй цену, если она не пришла из подключенного прайса. "
    "Если найденный прайс не содержит цену, переводи запрос мастеру. "
    "Не упоминай внутренние API, модели или служебные ошибки."
)

price_catalog = GoogleSheetsPriceCatalog.from_environment()


def price_request(text: str) -> bool:
    normalized = text.casefold()
    brands = ("iphone", "айфон", "samsung", "xiaomi", "redmi", "poco", "tecno", "infinix", "honor", "realme", "oppo")
    services = ("цена", "стоимость", "ремонт", "замена", "экран", "дисплей", "акб", "батаре", "заряд", "динамик", "кнопк")
    return any(item in normalized for item in brands) and any(item in normalized for item in services)


@dp.message(CommandStart())
async def start_cmd(message: types.Message):
    await message.answer(
        "Привет! Я Фаина 👋\n\n"
        "Я умею отвечать на русском и таджикском, распознавать текст и таблицы с фото, "
        "а также создавать Word и Excel.\n\n"
        "Команды:\n"
        "• /help — помощь\n"
        "• /word — создать Word из текста\n"
        "• /excel — создать Excel из текста\n\n"
        "Просто отправь сообщение или фотографию."
    )


@dp.message(Command("help"))
async def help_cmd(message: types.Message):
    await message.answer(
        "Фаина умеет:\n"
        "• 💬 отвечать на русском и таджикском\n"
        "• 📷 распознавать текст и таблицы с фото\n"
        "• 📄 создавать Word (.docx)\n"
        "• 📊 создавать Excel (.xlsx)\n\n"
        "Напиши /word или /excel, затем данные. Также можно написать «сделай Word» или «сделай Excel»."
    )


async def extract_photo(photo_bytes: bytes, caption: str = "", mime_type: str = "image/jpeg") -> dict[str, Any] | None:
    prompt = """
Ты выполняешь точное OCR и извлечение структуры документа с изображения.
Верни ТОЛЬКО валидный JSON без markdown:
{
  "title": "название документа или пусто",
  "language": "ru|tg|mixed|other",
  "document_text": "весь обычный текст в естественном порядке",
  "headers": ["название колонки 1", "название колонки 2"],
  "rows": [["значение", "значение"]]
}
Правила:
1. Сохраняй исходный текст максимально дословно.
2. Обязательно сохраняй таджикские буквы Ғ Ӣ Қ Ӯ Ҳ Ҷ и регистр.
3. Если на фото таблица, восстанови реальные названия колонок и все строки.
4. Не придумывай отсутствующие значения.
5. Если таблицы нет, headers и rows должны быть пустыми массивами.
6. document_text должен содержать весь распознанный текст, включая текст таблицы, если это необходимо для полноты.
""".strip()
    if caption:
        prompt += f"\nДополнительная инструкция пользователя: {caption[:2000]}"
    response = await ask_gemini(
        [genai_types.Part.from_bytes(data=photo_bytes, mime_type=mime_type), prompt],
        "Ты специалист по OCR документов и таблиц."
    )
    if not response:
        return None
    try:
        data = json.loads(clean_json(response))
        if not isinstance(data, dict):
            return None
        return normalize_data(data)
    except Exception as e:
        print(f"OCR JSON parse failed: {type(e).__name__}", flush=True)
        return normalize_data({
            "title": "Распознанный документ",
            "language": "mixed",
            "document_text": response,
            "headers": [],
            "rows": [],
        })


def make_docx(data: dict[str, Any]) -> bytes:
    data = normalize_data(data)
    doc = Document()
    title = data["title"]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(16)

    text = data["document_text"]
    if text:
        for block in re.split(r"\n\s*\n", text):
            if block.strip():
                doc.add_paragraph(block.strip())

    headers = data["headers"]
    rows = data["rows"]
    if headers:
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for row in rows:
            cells = table.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = value
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def make_xlsx(data: dict[str, Any]) -> bytes:
    data = normalize_data(data)
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    headers, rows, text = data["headers"], data["rows"], data["document_text"]
    if headers:
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in range(1, len(headers) + 1):
            max_len = max([len(str(ws.cell(r, col).value or "")) for r in range(1, ws.max_row + 1)] + [10])
            ws.column_dimensions[get_column_letter(col)].width = min(max_len + 2, 45)
    else:
        ws.append([data["title"]])
        for line in text.splitlines():
            if line.strip():
                ws.append([line.strip()])
        ws.column_dimensions["A"].width = min(max([len(str(c.value or "")) for c in ws["A"]] + [20]) + 2, 80)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


async def send_file_pair(message: types.Message, data: dict[str, Any], word: bool = True, excel: bool = True):
    data = normalize_data(data)
    title = re.sub(r"[^\w\- ]+", "", data["title"], flags=re.UNICODE).strip()[:50] or "Faina"
    if word:
        await message.answer_document(BufferedInputFile(make_docx(data), filename=f"{title}.docx"), caption="📄 Word готов")
    if excel:
        await message.answer_document(BufferedInputFile(make_xlsx(data), filename=f"{title}.xlsx"), caption="📊 Excel готов")


async def process_photo_message(message: types.Message, file_id: str, caption: str, mime_type: str = "image/jpeg"):
    async with lock_for(message.chat.id):
        try:
            await bot.send_chat_action(chat_id=message.chat.id, action="typing")
        except Exception:
            pass
        try:
            file = await bot.get_file(file_id)
            stream = io.BytesIO()
            await bot.download_file(file.file_path, stream)
            data = await extract_photo(stream.getvalue(), caption, mime_type)
            if not data:
                await message.answer("Не удалось распознать изображение. Попробуй отправить фото ещё раз.")
                return
            summary = data["document_text"]
            if summary:
                await message.answer("📷 Распознано:\n\n" + summary[:2500])
            await send_file_pair(message, data, word=True, excel=True)
        except Exception as e:
            print(f"Photo processing failed: {type(e).__name__}: {e}", flush=True)
            await message.answer("Не удалось обработать фото. Попробуй ещё раз.")


@dp.message(F.photo)
async def photo_handler(message: types.Message):
    photo = message.photo[-1]
    await process_photo_message(message, photo.file_id, message.caption or "", "image/jpeg")


@dp.message(F.document)
async def image_document_handler(message: types.Message):
    mime = message.document.mime_type or ""
    if not mime.startswith("image/"):
        return
    await process_photo_message(message, message.document.file_id, message.caption or "", mime)


async def generate_document_from_text(message: types.Message, text: str, word: bool, excel: bool):
    extraction_prompt = (
        "Преобразуй данные пользователя в JSON для документа. Верни только JSON: "
        '{"title":"...","language":"ru|tg|mixed","document_text":"...",'
        '"headers":["..."],"rows":[["..."]]}. '
        "Сохраняй таджикские буквы Ғ Ӣ Қ Ӯ Ҳ Ҷ. Если таблица не нужна, headers/rows пустые. "
        "Не выдумывай данные.\n\nДанные пользователя:\n" + text[:MAX_TEXT]
    )
    raw = await ask_gemini(extraction_prompt, "Ты специалист по структурированию документов.")
    if not raw:
        await message.answer("Сейчас не удалось создать документ. Попробуй ещё раз.")
        return
    try:
        data = normalize_data(json.loads(clean_json(raw)))
    except Exception:
        data = normalize_data({"title": "Документ", "document_text": raw, "headers": [], "rows": []})
    await send_file_pair(message, data, word=word, excel=excel)


@dp.message(Command("word"))
async def word_cmd(message: types.Message):
    text = (message.text or "").partition(" ")[2].strip()
    if not text:
        await message.answer("Напиши после команды данные, например: /word Справка о работе...")
        return
    async with lock_for(message.chat.id):
        await generate_document_from_text(message, text, word=True, excel=False)


@dp.message(Command("excel"))
async def excel_cmd(message: types.Message):
    text = (message.text or "").partition(" ")[2].strip()
    if not text:
        await message.answer("Напиши после команды данные, например: /excel Имя, Телефон...")
        return
    async with lock_for(message.chat.id):
        await generate_document_from_text(message, text, word=False, excel=True)


async def text_handler(message: types.Message):
    if not message.text:
        return
    text = message.text.strip()
    print(f"Telegram message received chat_id={message.chat.id} text_len={len(text)}", flush=True)
    lower = text.lower()
    wants_word = any(x in lower for x in ("сделай word", "создай word", "word документ", "/word"))
    wants_excel = any(x in lower for x in ("сделай excel", "создай excel", "excel таблиц", "/excel"))
    async with lock_for(message.chat.id):
        try:
            await bot.send_chat_action(chat_id=message.chat.id, action="typing")
        except Exception:
            pass
        if wants_word or wants_excel:
            await generate_document_from_text(message, text, wants_word, wants_excel)
            return
        catalog_context = ""
        if price_request(text):
            if not price_catalog:
                await message.answer("Контакты Мастера:\nДля точного расчёта стоимости свяжитесь напрямую с мастером.\n👉 Telegram: @MATIN_0893 @Coichi")
                return
            result = await price_catalog.lookup(text)
            if result.status in {"error", "not_found"}:
                await message.answer("Контакты Мастера:\nДля точного расчёта стоимости свяжитесь напрямую с мастером.\n👉 Telegram: @MATIN_0893 @Coichi")
                return
            catalog_context = "\nТочная строка прайса (используй только эти значения):\n" + json.dumps(result.row, ensure_ascii=False)
        prompt = SYSTEM_PROMPT + catalog_context + "\n\nСообщение пользователя:\n" + text[:MAX_TEXT]
        answer = await ask_gemini(prompt)
        if answer:
            for i in range(0, len(answer), 4000):
                await message.answer(answer[i:i + 4000])
        else:
            await message.answer("Сейчас AI временно занят. Попробуй ещё раз через несколько секунд.")


@dp.message(F.text)
async def text_message_handler(message: types.Message):
    await text_handler(message)


async def health_check(request: web.Request):
    return web.json_response({"status": "ok", "service": "faina-bot"})


async def configure_webhook():
    parsed_url = urlparse(BASE_URL)
    if parsed_url.scheme != "https" or not parsed_url.netloc:
        print("A valid HTTPS RENDER_EXTERNAL_URL or BASE_URL is required for webhook", flush=True)
        return False
    webhook_url = f"{BASE_URL}{WEBHOOK_PATH}"
    allowed_updates = dp.resolve_used_update_types()
    for attempt in range(1, 4):
        try:
            print(f"Setting Telegram webhook attempt={attempt} url={webhook_url}", flush=True)
            await asyncio.wait_for(
                bot.set_webhook(
                    url=webhook_url,
                    secret_token=WEBHOOK_SECRET or None,
                    drop_pending_updates=False,
                    max_connections=20,
                    allowed_updates=allowed_updates,
                ),
                timeout=12,
            )
            info = await asyncio.wait_for(bot.get_webhook_info(), timeout=8)
            print(f"Telegram webhook set successfully pending={info.pending_update_count} last_error={info.last_error_message!r}", flush=True)
            return True
        except Exception as e:
            print(f"Webhook setup failed attempt={attempt}: {type(e).__name__}: {e}", flush=True)
            await asyncio.sleep(attempt * 2)
    print("Webhook setup exhausted retries; server remains available", flush=True)
    return False


async def run_polling():
    print("Starting Telegram polling because no valid webhook URL is configured", flush=True)
    await bot.delete_webhook(drop_pending_updates=False)
    await dp.start_polling(bot, handle_signals=False)


async def run_runtime():
    if BASE_URL and urlparse(BASE_URL).scheme == "https" and urlparse(BASE_URL).netloc:
        if await configure_webhook():
            return
        print("Webhook setup failed; switching to polling", flush=True)
    await run_polling()


async def on_startup(app: web.Application):
    global runtime_task
    print(
        f"Faina bot starting on port={os.getenv('PORT', 'unknown')} "
        f"base_url={BASE_URL or '<not configured>'} models={','.join(MODELS)}",
        flush=True,
    )
    runtime_task = asyncio.create_task(run_runtime())
    print("Faina startup complete; Telegram update loop running in background", flush=True)


async def on_shutdown(app: web.Application):
    global runtime_task
    if runtime_task:
        runtime_task.cancel()
        try:
            await runtime_task
        except asyncio.CancelledError:
            pass
    try:
        await bot.delete_webhook(drop_pending_updates=False)
    except Exception as e:
        print(f"Webhook cleanup failed: {type(e).__name__}: {e}", flush=True)
    await bot.session.close()


app = web.Application()
app.router.add_get("/", health_check)
app.router.add_get("/health", health_check)

handler = SimpleRequestHandler(dispatcher=dp, bot=bot, secret_token=WEBHOOK_SECRET or None)
handler.register(app, path=WEBHOOK_PATH)
setup_application(app, dp, bot=bot)
app.on_startup.append(on_startup)
app.on_shutdown.append(on_shutdown)

if __name__ == "__main__":
    port = int(os.getenv("PORT", "10000"))
    print(f"Starting HTTP server on 0.0.0.0:{port}", flush=True)
    web.run_app(app, host="0.0.0.0", port=port)